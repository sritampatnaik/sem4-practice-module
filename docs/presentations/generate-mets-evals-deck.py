#!/usr/bin/env python3
"""Generate METS Evals presentation deck (PPTX, importable into Google Slides)."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

# METS palette (from architecture diagram)
METS_BLUE = RGBColor(0x42, 0x85, 0xF4)
METS_BLUE_DARK = RGBColor(0x1A, 0x56, 0xC4)
METS_GOLD = RGBColor(0xF9, 0xAB, 0x00)
METS_SLATE = RGBColor(0x3C, 0x40, 0x43)
METS_MUTED = RGBColor(0x5F, 0x63, 0x68)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_BG = RGBColor(0xE8, 0xF0, 0xFE)


def set_run(run, *, size=18, bold=False, color=METS_SLATE, font="Calibri"):
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def add_title_slide(prs: Presentation):
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank

    # Header band
    shape = slide.shapes.add_shape(
        1, Inches(0), Inches(0), prs.slide_width, Inches(2.4)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = METS_BLUE
    shape.line.fill.background()

    title_box = slide.shapes.add_textbox(Inches(0.7), Inches(0.55), Inches(8.6), Inches(1.2))
    tf = title_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    run = p.add_run()
    run.text = "METS Evals"
    set_run(run, size=40, bold=True, color=WHITE)

    p2 = tf.add_paragraph()
    run2 = p2.add_run()
    run2.text = "What they are & why they matter"
    set_run(run2, size=24, color=WHITE)

    subtitle = slide.shapes.add_textbox(Inches(0.7), Inches(2.9), Inches(8.6), Inches(1.5))
    stf = subtitle.text_frame
    stf.word_wrap = True
    sp = stf.paragraphs[0]
    sr = sp.add_run()
    sr.text = (
        "Multi-Agent Educational & Testing System\n"
        "Singapore Primary · O-Level · A-Level tutor"
    )
    set_run(sr, size=18, color=METS_MUTED)

    footer = slide.shapes.add_textbox(Inches(0.7), Inches(6.8), Inches(8.6), Inches(0.5))
    fp = footer.text_frame.paragraphs[0]
    fr = fp.add_run()
    fr.text = "METS Team · August 2026"
    set_run(fr, size=12, color=METS_MUTED)


def add_section_header(slide, slide_width, title: str, subtitle: str = ""):
    band = slide.shapes.add_shape(1, Inches(0), Inches(0), slide_width, Inches(1.05))
    band.fill.solid()
    band.fill.fore_color.rgb = METS_BLUE
    band.line.fill.background()

    tb = slide.shapes.add_textbox(Inches(0.55), Inches(0.18), Inches(8.8), Inches(0.7))
    tf = tb.text_frame
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = title
    set_run(r, size=28, bold=True, color=WHITE)

    if subtitle:
        sb = slide.shapes.add_textbox(Inches(0.55), Inches(1.25), Inches(8.8), Inches(0.45))
        sp = sb.text_frame.paragraphs[0]
        sr = sp.add_run()
        sr.text = subtitle
        set_run(sr, size=14, color=METS_MUTED)


def add_bullets(slide, items, top=1.75, left=0.65, width=8.7, size=17, spacing=0.52):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(5.5))
    tf = box.text_frame
    tf.word_wrap = True

    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(8)
        p.level = 0
        run = p.add_run()
        run.text = item
        set_run(run, size=size, color=METS_SLATE)


def add_callout(slide, text: str, top=6.0):
    box = slide.shapes.add_shape(1, Inches(0.55), Inches(top), Inches(8.9), Inches(0.75))
    box.fill.solid()
    box.fill.fore_color.rgb = LIGHT_BG
    box.line.color.rgb = METS_BLUE

    tb = slide.shapes.add_textbox(Inches(0.75), Inches(top + 0.12), Inches(8.5), Inches(0.55))
    p = tb.text_frame.paragraphs[0]
    r = p.add_run()
    r.text = text
    set_run(r, size=13, bold=True, color=METS_BLUE_DARK)


def add_content_slide(prs, title, subtitle, bullets, callout=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_section_header(slide, prs.slide_width, title, subtitle)
    add_bullets(slide, bullets)
    if callout:
        add_callout(slide, callout)


def build_deck() -> Presentation:
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)

    add_title_slide(prs)

    add_content_slide(
        prs,
        "What are evals?",
        "Systematic checks that our AI tutor behaves correctly",
        [
            "Evaluations (evals) are repeatable tests that measure whether METS agents "
            "do the right thing — not just whether they produce text.",
            "In METS, evals cover routing (Teaching vs Testing), subject selection, "
            "grade-band alignment, and assessment quality (MCQs & flashcards).",
            "Unlike ad-hoc manual chat tests, evals use fixed scenarios and profiles "
            "so we can compare results across prompt changes and model updates.",
            "Goal: catch regressions early and give educators confidence that outputs "
            "are syllabus-aligned, safe, and pedagogically sound.",
        ],
        callout="Evals answer: “Did the system behave correctly?” — not just “Did it reply?”",
    )

    add_content_slide(
        prs,
        "What we are building",
        "Foundations in place · planned eval layers next",
        [
            "Testing harness — POST /api/testing-harness runs 6 fixture scenarios "
            "(Primary fractions, O-Level kinematics, JC differentiation, etc.) in isolation.",
            "Observability — routing audit trail, prompt versioning, latency & token logs "
            "in logs/prompts.jsonl and the Studio UI routing panel.",
            "Schema validation — MCQ & flashcard tool contracts (IDs, option counts, "
            "correctOptionId) enforced before widgets render.",
            "Planned next — routing evals, RAG retrieval accuracy, verifier/marker for "
            "generated MCQs, and user-acceptance evaluation (Phase 4.2).",
        ],
        callout="Today: harness + observability. Tomorrow: automated quality gates.",
    )

    add_content_slide(
        prs,
        "Why evals matter for METS",
        "Trust, safety, and educational integrity in Singapore",
        [
            "Students depend on accurate explanations and fair assessments — evals guard "
            "against wrong answers, off-syllabus content, and inappropriate difficulty.",
            "Educators need explainability — every routing decision is logged with "
            "rationale and confidence so misroutes can be audited and fixed.",
            "Multi-agent systems fail silently — a bad router or weak Testing prompt "
            "can look fine in demo but break at scale; evals surface these early.",
            "Singapore governance alignment — MOE EdTech goals and AI assurance require "
            "traceability, fairness, and curriculum grounding, not black-box replies.",
        ],
        callout="Evals turn METS from a demo into a trustworthy educational product.",
    )

    add_content_slide(
        prs,
        "Roadmap & next steps",
        "How the team uses evals going forward",
        [
            "Run live harness scenarios once OPENAI_API_KEY is configured; record widget "
            "choice, tool outputs, and grade-band fit for each fixture.",
            "Add tool-contract checks — automated pass/fail on MCQ schema, flashcard "
            "structure, and performance-log writes.",
            "Extend evals to Orchestration routing (intent, subject, grade band) and "
            "RAG retrieval accuracy across Math, Physics, and Chemistry.",
            "Build verifier/marker module — validate generated MCQ correctness and "
            "track student outcomes before Phase 4 user-acceptance evaluation.",
        ],
        callout="Metric to watch: repeatable scenario pass rate + routing confidence.",
    )

    return prs


def main():
    out_dir = Path(__file__).resolve().parent
    out_path = out_dir / "METS-Evals-Deck.pptx"
    artifact_path = Path("/opt/cursor/artifacts/METS-Evals-Deck.pptx")

    prs = build_deck()
    prs.save(out_path)
    prs.save(artifact_path)

    print(f"Created: {out_path}")
    print(f"Copied:  {artifact_path}")
    print("\nTo open in Google Slides:")
    print("  1. Go to https://slides.google.com")
    print("  2. File → Import slides → Upload → select METS-Evals-Deck.pptx")


if __name__ == "__main__":
    main()
